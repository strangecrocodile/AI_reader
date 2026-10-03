"""阅读工具：翻译与总结。

测试重点不是「模型译得好不好」（那是模型的事），而是我们自己的那几层约定：

- SSE 事件形状与问答一致（meta → delta → done），前端才能复用同一个解析器；
- **依据由我们反查**：总结里只认 `[n]` 编号，锚点从编号查表得出，不信模型自报 id；
- 作用范围有界：翻译只处理选中的那段或那一页，绝不把整章几百段一起送出去；
- 无模型时不编造：翻译如实说「译不了」并列出原文，总结退化成规则摘要并标明来源；
- 章总结缓存，换章/重开仍在；教材被替换或删除时缓存必须作废。
"""
import io
import json

import pytest
from docx import Document

from app.services import tools as tools_service
from app.services.tools import TOOL_SUMMARY


def _docx(chapters):
    document = Document()
    for title, paragraphs in chapters:
        document.add_heading(title, level=1)
        for text in paragraphs:
            document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

BOOK = _docx(
    [
        (
            "第1章 函数与极限",
            [
                "极限描述的是当自变量无限接近某个值时，函数值趋近的那个数。",
                "函数是描述两个集合之间对应关系的工具。",
                "连续函数在闭区间上必定取得最大值和最小值。",
            ],
        ),
        ("第2章 导数与微分", ["导数刻画的是函数在一点处的变化率，几何意义是切线斜率。"]),
    ]
)


def _upload(client, data=BOOK, filename="微积分入门.docx"):
    res = client.post("/api/books", files={"file": (filename, data, DOCX_TYPE)})
    assert res.status_code == 201, res.text
    return res.json()


def _events(response):
    """把 SSE 响应体解析成事件列表（与前端解析器同一套约定）。"""
    out = []
    for block in response.text.split("\n\n"):
        for line in block.split("\n"):
            if line.startswith("data: "):
                out.append(json.loads(line[6:]))
    return out


def _done(events):
    finals = [e for e in events if e["event"] == "done"]
    assert finals, f"事件流没有 done 收尾：{events}"
    return finals[-1]


def _chapter(client, book, index=0):
    return book["chapters"][index]["id"]


# ---------------------------------------------------------------- 翻译


def test_translate_page_streams_each_paragraph_with_its_anchor(client):
    """「翻译本页」：逐段给出译文，每段带自己的锚点，前端才能左右对照。"""
    book = _upload(client)
    chapter_id = _chapter(client, book)

    res = client.post(
        "/api/tools/translate",
        json={"bookId": book["id"], "chapterId": chapter_id, "page": 1},
    )

    assert res.status_code == 200
    assert res.headers["content-type"].startswith("text/event-stream")
    events = _events(res)
    meta = next(e for e in events if e["event"] == "meta")
    assert meta["scope"] == "page"
    assert meta["target"] == "en", "中文教材的默认翻译方向是英文（见语言自动判定那条用例）"
    assert meta["paragraphs"], "这一页应有待译段落"
    assert all(p["anchorId"] for p in meta["paragraphs"]), "每段都要带锚点"

    done = _done(events)
    assert done["scope"] == "page"
    assert len(done["translations"]) == len(meta["paragraphs"])
    assert [t["anchorId"] for t in done["translations"]] == [p["anchorId"] for p in meta["paragraphs"]]


def test_translate_selection_uses_the_selected_text_only(client):
    book = _upload(client)
    chapter_id = _chapter(client, book)

    res = client.post(
        "/api/tools/translate",
        json={
            "bookId": book["id"],
            "chapterId": chapter_id,
            "selectedText": "函数是描述两个集合之间对应关系的工具。",
        },
    )
    done = _done(_events(res))

    assert done["scope"] == "selection"
    assert len(done["translations"]) == 1
    assert "对应关系" in done["translations"][0]["source"]


def test_translate_target_language_is_auto_but_can_be_forced(client, app):
    """中文原文默认译成英文；显式指定目标语言时听用户的。"""
    book = _upload(client)
    chapter_id = _chapter(client, book)

    auto = _done(
        _events(
            client.post(
                "/api/tools/translate",
                json={"bookId": book["id"], "chapterId": chapter_id, "page": 1},
            )
        )
    )
    assert auto["target"] == "en", "中文原文的默认方向是译成英文"

    forced = _done(
        _events(
            client.post(
                "/api/tools/translate",
                json={"bookId": book["id"], "chapterId": chapter_id, "page": 1, "target": "zh"},
            )
        )
    )
    assert forced["target"] == "zh"


def test_translate_english_source_defaults_to_chinese():
    assert tools_service.resolve_target_language("A limit describes the behaviour.") == "zh"
    assert tools_service.resolve_target_language("极限描述的是函数值趋近的那个数。") == "en"
    assert tools_service.resolve_target_language("abc 中文混排的句子") == "en"
    assert tools_service.resolve_target_language("any text", "zh") == "zh"


def test_translate_without_model_states_the_truth_instead_of_inventing(client):
    """无模型时必须说「译不了」并给出原文，而不是编一段像译文的中文。

    编一段是最坏的结果：用户会以为那就是原文的意思。
    """
    book = _upload(client)  # 测试环境用 MockLLM（kind=mock）
    chapter_id = _chapter(client, book)

    done = _done(
        _events(
            client.post(
                "/api/tools/translate",
                json={"bookId": book["id"], "chapterId": chapter_id, "page": 1},
            )
        )
    )

    assert done["notice"] == tools_service.FALLBACK_NOTICE
    assert all(t["text"] == "" for t in done["translations"]), "没有模型就不该有译文"
    assert all(t["source"] for t in done["translations"]), "原文要照常给出来对照"


def test_translate_page_scope_never_sends_the_whole_chapter(client, app):
    """作用范围必须有界：没给页码时退到本章首页，而不是整章几百段。"""
    book = _upload(client)
    chapter_id = _chapter(client, book)
    sections = app.state.db.sections_of(book["id"], chapter_id)
    pages = {s["page"] for s in sections if s["kind"] != "heading"}
    assert len(pages) >= 1

    done = _done(
        _events(client.post("/api/tools/translate", json={"bookId": book["id"], "chapterId": chapter_id}))
    )
    translated = {t["page"] for t in done["translations"]}
    assert len(translated) == 1, f"一次只该译一页，实际译了 {translated}"


def test_translate_unknown_book_or_chapter_is_404(client):
    book = _upload(client)
    assert (
        client.post(
            "/api/tools/translate", json={"bookId": "nope", "chapterId": "x"}
        ).status_code
        == 404
    )
    assert (
        client.post(
            "/api/tools/translate", json={"bookId": book["id"], "chapterId": "nope"}
        ).status_code
        == 404
    )


# ---------------------------------------------------------------- 总结


def test_summarize_chapter_streams_and_cites_anchors(client, app):
    book = _upload(client)
    chapter_id = _chapter(client, book)

    res = client.post(
        "/api/tools/summarize", json={"bookId": book["id"], "chapterId": chapter_id}
    )

    events = _events(res)
    meta = next(e for e in events if e["event"] == "meta")
    assert meta["scope"] == "chapter"
    assert meta["chunkCount"] >= 1
    assert meta["paragraphCount"] >= 1

    done = _done(events)
    assert done["summary"], "总结不能是空的"
    # 依据必须落回本章真实存在的锚点
    valid = {a["id"] for a in app.state.db.anchors_of(book["id"], chapter_id)}
    assert done["sources"], "总结要带依据"
    assert set(done["sources"]) <= valid
    assert all(d["id"] in valid for d in done["sourceDetails"])


def test_summarize_rule_fallback_is_labelled(client):
    """无模型时的规则摘要必须自报家门：它不是模型写的总结。"""
    book = _upload(client)
    chapter_id = _chapter(client, book)

    done = _done(
        _events(client.post("/api/tools/summarize", json={"bookId": book["id"], "chapterId": chapter_id}))
    )

    assert done["notice"] == tools_service.RULE_SUMMARY_NOTICE
    # 规则摘要的每一句都出自原文
    assert "极限描述的是" in done["summary"] or "函数是描述" in done["summary"]


def test_summarize_selection_scope(client):
    book = _upload(client)
    chapter_id = _chapter(client, book)

    done = _done(
        _events(
            client.post(
                "/api/tools/summarize",
                json={
                    "bookId": book["id"],
                    "chapterId": chapter_id,
                    "selectedText": "连续函数在闭区间上必定取得最大值和最小值。",
                },
            )
        )
    )
    assert done["scope"] == "selection"


def test_chapter_summary_is_cached_and_readable(client, app):
    """章总结要缓存：重开这一章不该再花一次模型调用（也不该等第二遍）。"""
    book = _upload(client)
    chapter_id = _chapter(client, book)

    before = client.get(f"/api/books/{book['id']}/chapters/{chapter_id}/summary").json()
    assert before == {"summary": "", "sources": [], "sourceDetails": [], "cached": False}

    done = _done(
        _events(client.post("/api/tools/summarize", json={"bookId": book["id"], "chapterId": chapter_id}))
    )

    after = client.get(f"/api/books/{book['id']}/chapters/{chapter_id}/summary").json()
    assert after["cached"] is True
    assert after["summary"] == done["summary"]
    assert after["sources"] == done["sources"]
    assert app.state.db.get_tool_output(book["id"], chapter_id, TOOL_SUMMARY)


def test_selection_summary_is_not_cached(client, app):
    """选中段落的总结不进缓存：它只对那一句话有意义，缓存下来会顶掉章总结。"""
    book = _upload(client)
    chapter_id = _chapter(client, book)

    _events(
        client.post(
            "/api/tools/summarize",
            json={"bookId": book["id"], "chapterId": chapter_id, "selectedText": "函数是描述两个集合之间对应关系的工具。"},
        )
    )

    assert app.state.db.get_tool_output(book["id"], chapter_id, TOOL_SUMMARY) is None


def test_summary_cache_dies_with_the_book_and_with_a_replace(client, app):
    """缓存挂在章节上：换掉教材内容或删掉教材，旧总结都必须消失。"""
    book = _upload(client)
    chapter_id = _chapter(client, book)
    _events(client.post("/api/tools/summarize", json={"bookId": book["id"], "chapterId": chapter_id}))
    assert app.state.db.get_tool_output(book["id"], chapter_id, TOOL_SUMMARY) is not None

    res = client.post(
        f"/api/books/{book['id']}/replace",
        files={"file": ("微积分入门.docx", BOOK, DOCX_TYPE)},
    )
    assert res.status_code == 200, res.text
    assert app.state.db.get_tool_output(book["id"], chapter_id, TOOL_SUMMARY) is None, (
        "教材内容换过了，基于旧段落的总结必须作废"
    )


def test_summary_cache_removed_on_book_delete(client, app):
    book = _upload(client)
    chapter_id = _chapter(client, book)
    _events(client.post("/api/tools/summarize", json={"bookId": book["id"], "chapterId": chapter_id}))

    assert client.delete(f"/api/books/{book['id']}").status_code == 204

    assert app.state.db.get_tool_output(book["id"], chapter_id, TOOL_SUMMARY) is None


def test_summarize_unknown_chapter_is_404(client):
    book = _upload(client)
    assert (
        client.post(
            "/api/tools/summarize", json={"bookId": book["id"], "chapterId": "nope"}
        ).status_code
        == 404
    )


# ---------------------------------------------------------------- 纯函数


def test_cited_anchors_ignores_made_up_ids():
    """模型只能报编号：越界编号、重复编号都忽略，一个都没有时才退回前三条。"""
    mapping = [
        {"index": 1, "anchor_id": "a1"},
        {"index": 2, "anchor_id": "a2"},
        {"index": 3, "anchor_id": "a3"},
        {"index": 4, "anchor_id": "a4"},
    ]
    assert tools_service.cited_anchors("概述 [2]，要点 [1][2][99]", mapping) == ["a2", "a1"]
    assert tools_service.cited_anchors("没有引用的总结", mapping) == ["a1", "a2", "a3"]


def test_chunks_of_respects_paragraph_and_char_budgets():
    paragraphs = [{"id": f"s{i}", "text": "字" * 100} for i in range(20)]
    chunks = tools_service.chunks_of(paragraphs, size=4, budget=1000)
    assert [len(c) for c in chunks] == [4, 4, 4, 4, 4]

    # 单段超过字符预算时独占一块，不把一段话切开
    big = [{"id": "s1", "text": "字" * 5000}, {"id": "s2", "text": "短"}]
    chunks = tools_service.chunks_of(big, size=4, budget=1000)
    assert [[p["id"] for p in c] for c in chunks] == [["s1"], ["s2"]]


def test_rule_summary_keeps_first_sentences_only():
    summary = tools_service.rule_summary(
        [{"text": "第一句。第二句。"}, {"text": "另一段的开头，没有句号"}]
    )
    lines = summary.splitlines()
    assert lines[0].startswith("本章共 2 段")
    assert lines[1] == "- 第一句。"
    assert lines[2].startswith("- 另一段的开头")


def test_paragraphs_of_skips_headings_and_other_pages():
    sections = [
        {"id": "s1", "kind": "heading", "text": "1.1 小节", "page": 3},
        {"id": "s2", "kind": "p", "text": "第三页正文。", "page": 3},
        {"id": "s3", "kind": "p", "text": "第四页正文。", "page": 4},
        {"id": "s4", "kind": "p", "text": "   ", "page": 4},
    ]
    assert [s["id"] for s in tools_service.paragraphs_of(sections)] == ["s2", "s3"]
    assert [s["id"] for s in tools_service.paragraphs_of(sections, 4)] == ["s3"]
