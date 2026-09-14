"""pytest 公共夹具：应用、客户端与示例教材。"""
import sys
from pathlib import Path

import pytest

BACKEND_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_DIR))
sys.path.insert(0, str(BACKEND_DIR / "scripts"))

from app.llm.client import MockLLM  # noqa: E402
from app.main import create_app  # noqa: E402


@pytest.fixture
def app(tmp_path):
    return create_app(db_path=tmp_path / "test.db", llm=MockLLM())


@pytest.fixture
def client(app):
    from fastapi.testclient import TestClient

    return TestClient(app)


@pytest.fixture(scope="session")
def demo_pdf_bytes(tmp_path_factory):
    import make_demo_pdf

    out = tmp_path_factory.mktemp("pdf") / "demo.pdf"
    make_demo_pdf.build_demo_pdf(out)
    return out.read_bytes()


@pytest.fixture(scope="session")
def demo_docx_bytes():
    """用 python-docx 现场生成的原创 Word 教材（内容团队原创，版权安全）。"""
    import io

    from docx import Document

    document = Document()
    document.add_heading("微积分入门（Word 版）", level=0)  # Title 样式 = 书名
    document.add_heading("第1章 函数与极限", level=1)
    document.add_heading("1.1 函数的概念", level=2)
    document.add_paragraph("数学上，我们把这种对应关系称为函数：对每个 x 都有唯一的 y 与之对应。")
    document.add_paragraph("函数的表示方法常见的有解析式、表格和图像三种。")
    document.add_heading("第2章 导数与微分", level=1)
    document.add_heading("2.1 导数的概念", level=2)
    document.add_paragraph("比值 Δy / Δx 衡量区间内的平均快慢，称为平均变化率。")
    document.add_paragraph("让 Δx 趋近于 0，平均变化率趋近的那个数就定义为导数。")
    document.add_paragraph("f(x) = 2x + 1")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture(scope="session")
def demo_markdown_bytes():
    """Markdown 教材样例：一级标题是书名，二级是章，三级是节。"""
    text = (
        "# 微积分入门（Markdown 版）\n\n"
        "## 第1章 函数与极限\n\n"
        "### 1.1 函数的概念\n\n"
        "数学上，我们把这种对应关系称为函数：对每个 x 都有唯一的 y 与之对应。\n\n"
        "## 第2章 导数与微分\n\n"
        "### 2.1 导数的概念\n\n"
        "让 Δx 趋近于 0，平均变化率趋近的那个数就定义为导数，记作 f′(x₀)。\n\n"
        "## 第3章 微分中值定理\n\n"
        "### 3.1 罗尔定理\n\n"
        "如果函数在闭区间上连续、在开区间内可导，就称它满足罗尔定理的条件。\n"
    )
    return text.encode("utf-8")


@pytest.fixture(scope="session")
def demo_docx_with_assets_bytes():
    """带**表格与插图**的 Word 教材样例（现场生成，版权安全）。

    这两类内容都不在 `doc.paragraphs` 里：表格是 body 的同级 `w:tbl`，
    图片藏在 run 的 `w:drawing` 里。所以专门造一份来验证它们确实被读出来了。
    """
    import io

    from docx import Document
    from docx.shared import Inches

    document = Document()
    document.add_heading("第1章 带图表的教材", level=1)
    document.add_paragraph("下面是一张表：")
    table = document.add_table(rows=2, cols=3)
    for index, text in enumerate(("概念", "记号", "含义")):
        table.cell(0, index).text = text
    for index, text in enumerate(("导数", "f′(x)", "瞬时变化率")):
        table.cell(1, index).text = text
    document.add_paragraph("下面是一张图：")
    document.add_picture(io.BytesIO(_tiny_png()), width=Inches(1.2))
    document.add_paragraph("图后还有一段正文。")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _tiny_png() -> bytes:
    """8x8 的纯色 PNG（够小，但足以验证落盘与可读）。"""
    import struct
    import zlib

    width = height = 8
    raw = b"".join(b"\x00" + bytes([200, 60, 60]) * width for _ in range(height))

    def chunk(tag: bytes, payload: bytes) -> bytes:
        return (
            struct.pack(">I", len(payload))
            + tag
            + payload
            + struct.pack(">I", zlib.crc32(tag + payload) & 0xFFFFFFFF)
        )

    header = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", header)
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


@pytest.fixture(scope="session")
def demo_pdf_with_image_bytes():
    """带插图的 PDF 样例（现场生成）：验证插图会被抽出、落盘、可访问。"""
    import io

    import pymupdf as fitz

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 80), "Chapter 1 Demo", fontsize=20)
    # 用英文句号断句（切段规则见 pdf._ends_sentence）：
    # PDF 内置字体画不出中文，样例用 ASCII 更稳妥。
    page.insert_text((72, 120), "Body text before the figure.", fontsize=11)
    page.insert_image(fitz.Rect(72, 160, 272, 300), stream=_tiny_png())
    page.insert_text((72, 340), "Body text after the figure.", fontsize=11)
    data = doc.tobytes()
    doc.close()
    return data
