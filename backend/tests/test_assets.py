"""插图与表格：从原文件抽出来、落盘、经接口下发并可直接访问。

覆盖：
- Word：表格抽成行列结构（首行当表头）且文字进纯文本，供检索与锚点用；
- Word：内嵌图片落盘，且**插在正确位置**（表格/图片前后的正文顺序不乱）；
- PDF：页内插图落盘，并按纵向位置插回正文顺序；
- 接口：章节内容下发 image / table 段落，插图经 /assets 可访问；
- 大书按页数阈值跳过表格提取，并如实提示用户。
"""
import io
from pathlib import Path

import pymupdf as fitz
from docx import Document
from docx.shared import Inches

from app.parsing.assets import TABLE_SCAN_MAX_PAGES
from app.parsing.docx import parse_docx_bytes
from app.parsing.pdf import parse_pdf_stream


def _sections(book):
    return [s for chapter in book.chapters for s in chapter.sections]


def _kinds(book):
    return [s.kind for s in _sections(book)]


# ---------- Word ----------


def test_docx_table_becomes_table_section_with_rows(demo_docx_with_assets_bytes):
    book = parse_docx_bytes(demo_docx_with_assets_bytes)
    tables = [s for s in _sections(book) if s.kind == "table"]

    assert len(tables) == 1
    table = tables[0]
    assert table.content["header"] is True
    assert [cell["text"] for cell in table.content["rows"][0]] == ["概念", "记号", "含义"]
    assert [cell["text"] for cell in table.content["rows"][1]] == ["导数", "f′(x)", "瞬时变化率"]
    # 纯文本版供检索用，单元格之间要有分隔，否则两列会黏成一个词
    assert "概念 | 记号 | 含义" in table.text


def test_docx_table_keeps_surrounding_paragraph_order(demo_docx_with_assets_bytes):
    """表格必须落在它前后两段文字之间，而不是被挪到章末。"""
    sections = _sections(parse_docx_bytes(demo_docx_with_assets_bytes))
    texts = [(s.kind, s.text) for s in sections]

    kinds = [kind for kind, _ in texts]
    before = next(i for i, (kind, text) in enumerate(texts) if "下面是一张表" in text)
    table_at = kinds.index("table")
    after = next(
        i for i, (kind, text) in enumerate(texts) if "图后还有一段正文" in text
    )

    assert before < table_at < after


def test_docx_table_widens_content_warning_but_not_false_alarm(demo_docx_with_assets_bytes):
    """表格能被读到之后，不能再报「跳过了表格」。"""
    book = parse_docx_bytes(demo_docx_with_assets_bytes)

    assert all("表格" not in note for note in book.notes)


def test_docx_images_are_written_to_assets(tmp_path, demo_docx_with_assets_bytes):
    book = parse_docx_bytes(demo_docx_with_assets_bytes, assets_dir=tmp_path, book_id="bk1")
    images = [s for s in _sections(book) if s.kind == "image"]

    assert len(images) == 1, "样例里的 1 张图应当被抽出来"
    asset = images[0].content["asset"]
    stored = Path(tmp_path) / asset
    assert stored.is_file()
    # 落盘的应当是**原始图片字节**，不是重新编码过的近似物
    assert stored.read_bytes().startswith(b"\x89PNG")
    assert asset.startswith("bk1-")


def test_docx_image_sits_between_its_neighbouring_paragraphs(tmp_path, demo_docx_with_assets_bytes):
    sections = _sections(
        parse_docx_bytes(demo_docx_with_assets_bytes, assets_dir=tmp_path, book_id="bk1")
    )
    kinds = [s.kind for s in sections]

    before = next(i for i, s in enumerate(sections) if "下面是一张图" in s.text)
    image_at = kinds.index("image")
    after = next(i for i, s in enumerate(sections) if "图后还有一段正文" in s.text)

    assert before < image_at < after


# ---------- PDF ----------


def test_pdf_image_is_written_and_ordered_by_position(tmp_path, demo_pdf_with_image_bytes):
    book = parse_pdf_stream(
        demo_pdf_with_image_bytes, "demo", assets_dir=tmp_path, book_id="bk2"
    )
    sections = _sections(book)
    images = [s for s in sections if s.kind == "image"]

    assert len(images) == 1
    stored = Path(tmp_path) / images[0].content["asset"]
    assert stored.is_file() and stored.stat().st_size > 0

    # 图在「图前」「图后」两段正文之间（正文用英文句号断句，见 _ends_sentence）
    idx_before = next(i for i, s in enumerate(sections) if "Body text before" in s.text)
    idx_image = sections.index(images[0])
    idx_after = next(i for i, s in enumerate(sections) if "Body text after" in s.text)
    assert idx_before < idx_image < idx_after


def test_pdf_asset_names_are_scoped_to_book(tmp_path, demo_pdf_with_image_bytes):
    book = parse_pdf_stream(
        demo_pdf_with_image_bytes, "demo", assets_dir=tmp_path, book_id="abc123"
    )
    image = next(s for s in _sections(book) if s.kind == "image")

    assert image.content["asset"].startswith("abc123-p1-")


def test_pdf_skips_table_extraction_over_page_limit(tmp_path, monkeypatch):
    """超过页数阈值时只提图、不提表格，并把这件事写进 notes。"""
    doc = fitz.open()
    for _ in range(TABLE_SCAN_MAX_PAGES + 1):
        page = doc.new_page()
        page.insert_text((72, 80), "Chapter 1 Demo", fontsize=20)
        page.insert_text((72, 120), "Body text for this page.", fontsize=11)
    data = doc.tobytes()
    doc.close()

    called = {"tables": 0}
    import app.parsing.pdf as pdf_module

    real = pdf_module.extract_page_assets

    def spy(page, page_no, book_id, assets_dir, **kwargs):
        if kwargs.get("with_tables"):
            called["tables"] += 1
        return real(page, page_no, book_id, assets_dir, **kwargs)

    monkeypatch.setattr(pdf_module, "extract_page_assets", spy)
    book = parse_pdf_stream(data, "demo", assets_dir=tmp_path, book_id="big")

    assert called["tables"] == 0, "超限的书不该去跑昂贵的表格提取"
    assert any("没有提取表格" in note for note in book.notes)


# ---------- 接口 ----------


def _upload(client, content: bytes, filename: str, content_type: str):
    res = client.post(
        "/api/books",
        files={"file": (filename, content, content_type)},
        data={"title": "图文教材"},
    )
    assert res.status_code == 201, res.text
    return res.json()


def test_chapter_content_exposes_image_and_table_paragraphs(client, demo_docx_with_assets_bytes):
    book = _upload(
        client,
        demo_docx_with_assets_bytes,
        "figures.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    chapter_id = book["chapters"][0]["id"]

    content = client.get(f"/api/books/{book['id']}/chapters/{chapter_id}").json()
    types = [para["type"] for para in content["paragraphs"]]

    assert "table" in types
    assert "image" in types

    table = next(p for p in content["paragraphs"] if p["type"] == "table")
    assert table["header"] is True
    assert table["rows"][0] == ["概念", "记号", "含义"]

    image = next(p for p in content["paragraphs"] if p["type"] == "image")
    assert image["src"].startswith("/assets/")
    # 插图是要被 <img> 直接取的，必须真的能访问到
    fetched = client.get(image["src"])
    assert fetched.status_code == 200
    assert fetched.content.startswith(b"\x89PNG")


def test_chapter_content_returns_styled_runs_for_docx(client):
    """Word 的粗体/上下标要一路走到接口层，前端才渲染得出来。"""
    document = Document()
    document.add_heading("第1章 富文本", level=1)
    paragraph = document.add_paragraph()
    paragraph.add_run("设 ")
    bold = paragraph.add_run("f(x)")
    bold.bold = True
    paragraph.add_run(" 的导数为 ")
    sub = paragraph.add_run("0")
    sub.font.subscript = True
    paragraph.add_run("。")
    buffer = io.BytesIO()
    document.save(buffer)

    book = _upload(
        client,
        buffer.getvalue(),
        "rich.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    content = client.get(f"/api/books/{book['id']}/chapters/{book['chapters'][0]['id']}").json()

    para = next(p for p in content["paragraphs"] if p["type"] == "p")
    src = para["segs"][0]
    # 锚点仍在最外层，锚点定位与跨章跳转的既有逻辑不受影响
    assert src["t"] == "src" and src["id"]
    # 有版式时锚点自身不再重复带正文（v 留空），正文只出现在子片段里
    assert src["v"] == ""
    styles = {seg["v"]: seg.get("style") for seg in src["segs"]}
    assert styles["f(x)"] == ["b"]
    assert styles["0"] == ["sub"]
    # 片段拼起来就是这一段的原文
    assert "".join(seg["v"] for seg in src["segs"]) == "设 f(x) 的导数为 0。"


def test_chapter_content_falls_back_to_plain_text_without_content(client):
    """没有任何版式信息的来源要给出和老版本一样的单片段结构。

    纯文本教材（.txt）不会产出 content，正好代表「无版式」这条路径；
    PDF/Word 现在大多会带 content，所以不能拿它们来测回退。
    """
    res = client.post(
        "/api/books",
        files={"file": ("plain.txt", "第1章 纯文本\n这是一段没有任何格式的正文。".encode("utf-8"), "text/plain")},
        data={"title": "纯文本教材"},
    )
    assert res.status_code == 201, res.text
    book = res.json()
    content = client.get(f"/api/books/{book['id']}/chapters/{book['chapters'][0]['id']}").json()

    para = next(p for p in content["paragraphs"] if p["type"] == "p")
    assert len(para["segs"]) == 1
    src = para["segs"][0]
    assert src["t"] == "src"
    # 无版式时正文直接放在锚点上，不再多包一层 segs——前端行为与升级前完全一致
    assert "segs" not in src
    assert src["v"] == "这是一段没有任何格式的正文。"


def test_assets_are_publicly_served(client, app, demo_docx_with_assets_bytes):
    """插图经 /assets 静态目录提供；校验落盘目录确实挂在了这个前缀上。"""
    book = _upload(
        client,
        demo_docx_with_assets_bytes,
        "figures.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    content = client.get(f"/api/books/{book['id']}/chapters/{book['chapters'][0]['id']}").json()
    src = next(p["src"] for p in content["paragraphs"] if p["type"] == "image")

    assert (Path(app.state.settings.assets_dir) / Path(src).name).is_file()
    assert client.get(src).status_code == 200
