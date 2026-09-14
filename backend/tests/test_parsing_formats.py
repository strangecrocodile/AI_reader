"""多格式解析测试：Word（.docx）与纯文本 / Markdown → 统一的 ParsedBook。

覆盖：Word 大纲层级识别章与节、无标题样式时退回文本模式、前言并入首章、
公式行标记、Markdown 层级、代码块保留、GBK 编码、格式判定与不支持格式的报错。
"""
import io

import pytest
from docx import Document

from app.parsing.base import CHARS_PER_PAGE, Block, assemble_book, build_blocks, chapter_level_of
from app.parsing.docx import parse_docx_bytes
from app.parsing.text import decode_bytes, parse_text_bytes, read_rows
from app.services.ingest import (
    DOCX,
    LOW_CONTENT_CHARS,
    PDF,
    TEXT,
    content_warning_of,
    detect_format,
    parse_bytes,
)


def _docx(paragraphs) -> bytes:
    """按 [(样式, 文本)] 生成 docx（样式为 'Heading 1' / 'Title' / 'Normal'）。"""
    document = Document()
    for style, text in paragraphs:
        if style == "Title":
            document.add_heading(text, level=0)
        elif style.startswith("Heading "):
            document.add_heading(text, level=int(style.split()[1]))
        else:
            document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _save(document) -> bytes:
    """把已组装的 Document 存成 bytes（需要表格/图片时用它，`_docx` 只加段落）。"""
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ---------- Word ----------


def test_parse_docx_recognizes_chapters_sections_and_title(demo_docx_bytes):
    book = parse_docx_bytes(demo_docx_bytes, "兜底书名")

    assert book.title == "微积分入门（Word 版）"
    assert [c.title for c in book.chapters] == ["第1章 函数与极限", "第2章 导数与微分"]

    first = book.chapters[0]
    kinds = [(s.kind, s.text) for s in first.sections]
    assert ("heading", "1.1 函数的概念") in kinds
    assert any(kind == "p" and "称为函数" in text for kind, text in kinds)
    # 标题不进正文全文（避免污染检索与讲解）
    assert "1.1 函数的概念" not in first.full_text
    # 虚拟页码：递增且落在章节区间内
    assert all(s.page >= 1 for s in first.sections)
    assert first.page_start <= first.sections[0].page <= first.page_end


def test_parse_docx_marks_formula_lines(demo_docx_bytes):
    book = parse_docx_bytes(demo_docx_bytes, "兜底书名")
    chapter = book.chapters[1]

    assert any(s.kind == "formula" and "2x + 1" in s.text for s in chapter.sections)


def test_parse_docx_without_heading_styles_falls_back_to_text_patterns():
    data = _docx(
        [
            ("Normal", "第1章 函数与极限"),
            ("Normal", "极限是微积分中第一个重要的工具。"),
            ("Normal", "第2章 导数与微分"),
            ("Normal", "导数刻画的是一个量相对于另一个量的变化率。"),
            ("Normal", "第3章 微分中值定理"),
            ("Normal", "本节讨论中值定理的三个基本结论。"),
        ]
    )

    book = parse_docx_bytes(data, "兜底书名")

    assert [c.title for c in book.chapters] == ["第1章 函数与极限", "第2章 导数与微分", "第3章 微分中值定理"]
    assert "极限是微积分中第一个重要的工具。" in book.chapters[0].full_text


def test_parse_docx_without_any_structure_becomes_single_chapter():
    data = _docx([("Normal", "这是一段没有章节结构的正文。"), ("Normal", "第二段正文。")])

    book = parse_docx_bytes(data, "兜底书名")

    assert len(book.chapters) == 1
    assert book.chapters[0].title == "兜底书名"
    assert len(book.chapters[0].sections) == 2


def test_parse_docx_front_matter_merges_into_first_chapter():
    data = _docx(
        [
            ("Normal", "本书是团队原创的演示教材，供教学演示使用。"),
            ("Heading 1", "第1章 函数与极限"),
            ("Normal", "极限描述的是趋势。"),
        ]
    )

    book = parse_docx_bytes(data, "兜底书名")

    assert [c.title for c in book.chapters] == ["第1章 函数与极限"]
    assert "本书是团队原创的演示教材" in book.chapters[0].full_text


def test_parse_docx_skips_heading_levels_above_chapter_level():
    """比「章」更浅的标题（篇名）不应变成章节，也不应混进正文段落。"""
    data = _docx(
        [
            ("Heading 1", "第一篇 基础"),
            ("Heading 2", "第1章 函数与极限"),
            ("Normal", "正文内容。"),
            ("Heading 2", "第2章 导数与微分"),
            ("Normal", "更多正文。"),
        ]
    )

    book = parse_docx_bytes(data, "兜底书名")

    assert [c.title for c in book.chapters] == ["第1章 函数与极限", "第2章 导数与微分"]
    assert all("第一篇" not in s.text for c in book.chapters for s in c.sections)


# ---------- Markdown / 纯文本 ----------


def test_parse_markdown_uses_heading_levels(demo_markdown_bytes):
    book = parse_text_bytes(demo_markdown_bytes, "兜底书名")

    assert book.title == "微积分入门（Markdown 版）"
    assert [c.title for c in book.chapters] == [
        "第1章 函数与极限",
        "第2章 导数与微分",
        "第3章 微分中值定理",
    ]
    first = book.chapters[0]
    assert ("heading", "1.1 函数的概念") in [(s.kind, s.text) for s in first.sections]
    assert "称为函数" in first.full_text


def test_parse_markdown_keeps_code_block_as_body():
    text = (
        "# 动手学 Python\n\n"
        "## 第1章 起步\n\n"
        "下面是一段示例代码：\n\n"
        "```python\n"
        "# 这是注释，不是 Markdown 标题\n"
        "print(range(5))\n"
        "```\n\n"
        "运行后会输出一串数字。\n"
    ).encode("utf-8")

    book = parse_text_bytes(text, "兜底书名")

    assert [c.title for c in book.chapters] == ["第1章 起步"]
    body = book.chapters[0].full_text
    assert "print(range(5))" in body
    assert "# 这是注释" in body  # 围栏内的 # 不当标题
    assert all(s.kind != "heading" for s in book.chapters[0].sections)


def test_parse_txt_recognizes_chinese_chapter_and_section_patterns():
    text = (
        "第1章 函数与极限\n"
        "第1节 函数的概念\n"
        "数学上，我们把这种对应关系称为函数。\n"
        "第2章 导数与微分\n"
        "让 Δx 趋近于 0，定义出导数。\n"
    ).encode("utf-8")

    book = parse_text_bytes(text, "兜底书名")

    assert [c.title for c in book.chapters] == ["第1章 函数与极限", "第2章 导数与微分"]
    assert ("heading", "第1节 函数的概念") in [(s.kind, s.text) for s in book.chapters[0].sections]


def test_parse_txt_decodes_gbk_encoding():
    raw = "第1章 函数与极限\n极限是微积分的基础。".encode("gb18030")

    book = parse_text_bytes(raw, "兜底书名")

    assert book.chapters[0].title == "第1章 函数与极限"
    assert "极限是微积分的基础。" in book.chapters[0].full_text
    assert decode_bytes(raw).startswith("第1章")


def test_parse_txt_without_structure_becomes_single_chapter():
    book = parse_text_bytes("只有一段正文，没有任何章节标记。".encode("utf-8"), "兜底书名")

    assert len(book.chapters) == 1
    assert book.chapters[0].title == "兜底书名"


def test_read_rows_ignores_code_fence_markers():
    rows = read_rows("# 书名\n```\n## 不是标题\n```\n正文。\n")

    assert rows == [("书名", 1), ("## 不是标题", 0), ("正文。", 0)]


# ---------- 装配与格式判定 ----------


def test_assemble_book_numbers_chapters_and_assigns_virtual_pages():
    long_text = "字" * (CHARS_PER_PAGE + 10)
    book = assemble_book(
        "测试教材",
        [
            Block(title="第1章 甲", paragraphs=[(long_text, "p")]),
            Block(title="第2章 乙", paragraphs=[("短段落。", "p")]),
            Block(title="第3章 空", paragraphs=[]),
        ],
    )

    assert [c.num for c in book.chapters] == [1, 2]
    assert [c.title for c in book.chapters] == ["第1章 甲", "第2章 乙"]
    # 超过一页字符预算后，后续章节的虚拟页码递增
    assert book.chapters[0].page_start == 1
    assert book.chapters[1].page_start == 2


def test_chapter_level_prefers_deeper_level_when_shallow_is_book_title():
    rows = [("微积分入门", 1), ("第1章 函数", 2), ("1.1 函数", 3), ("第2章 导数", 2)]

    assert chapter_level_of(rows) == 2
    blocks = build_blocks(rows)
    assert [b.title for b in blocks] == ["第1章 函数", "第2章 导数"]
    # Block.paragraphs 是 (文本, kind, 富文本内容) 三元组；纯文本来源的 content 为
    # None，渲染层据此把文本当成单个无样式片段（见 base.runs_of_content）
    assert ("1.1 函数", "heading", None) in blocks[0].paragraphs


def test_chapter_level_falls_back_to_text_mode_for_title_only_headings():
    rows = [("微积分入门", 1)] + [(f"第{i}章 标题{i}", 0) for i in range(1, 5)]

    assert chapter_level_of(rows) is None
    assert len(build_blocks(rows)) == 4


def test_detect_format_by_extension_and_content_type():
    assert detect_format("book.pdf") == PDF
    assert detect_format("BOOK.DOCX") == DOCX
    assert detect_format("notes.md") == TEXT
    assert detect_format("notes.markdown") == TEXT
    assert detect_format("notes.txt") == TEXT
    # 扩展名不可用时退回 content-type
    assert detect_format("", "application/pdf") == PDF
    assert detect_format("blob", "text/plain") == TEXT
    assert detect_format(
        "blob", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ) == DOCX
    # 不支持的格式
    assert detect_format("scan.png", "image/png") is None
    assert detect_format("old.doc", "application/msword") is None


def test_parse_bytes_rejects_unknown_format():
    with pytest.raises(ValueError, match="docx"):
        parse_bytes("nope", b"data", "书名")


# ---------- 解析受限提示（content_warning 的根因信号） ----------


def test_clean_docx_reports_no_skipped_content():
    """全部内容都在正文段落里 → 不报任何「跳过了什么」。"""
    book = parse_docx_bytes(_docx([("Heading 1", "第1章 测试"), ("Normal", "正文段落。")]))

    assert book.notes == []


def test_docx_reports_skipped_table():
    """表格里的文字 `doc.paragraphs` 读不到，必须如实报告数量。"""
    document = Document()
    document.add_heading("第1章 测试", level=1)
    document.add_paragraph("正文段落。")
    document.add_table(rows=2, cols=2)

    book = parse_docx_bytes(_save(document))

    assert book.notes == ["1 个表格"]
    # 报告的数字要和实际丢失的内容对得上：表格文字确实一个字都没进正文
    assert all("表格" not in chapter.full_text for chapter in book.chapters)


def test_content_warning_flags_book_whose_text_hides_in_tables():
    """用户实际遇到的那个情况：正文都在表格里，整本只解析出几十个字。"""
    document = Document()
    document.add_heading("第1章 测试", level=1)
    document.add_paragraph("短短一句话。")
    document.add_table(rows=3, cols=2)

    warning = content_warning_of(parse_docx_bytes(_save(document)))

    assert "6 个字" in warning  # 症状：用户能直接感知的
    assert "1 个表格" in warning  # 根因：解析器跳过了什么


def test_content_warning_does_not_fire_on_thin_but_valid_textbook():
    """薄教材是合法的：kb-agent 样例书正文只有 523 字，不能被误报。"""
    body = "教" * (LOW_CONTENT_CHARS + 23)  # 与样例书正文同量级
    book = parse_docx_bytes(_docx([("Heading 1", "第1章 测试"), ("Normal", body)]))

    assert content_warning_of(book) == ""
