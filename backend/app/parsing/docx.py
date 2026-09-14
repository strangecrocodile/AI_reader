"""Word（.docx）解析：按 Word 大纲层级识别章与节，产出统一的 ParsedBook。

- 用 Word 原生标题样式（Heading 1..9）判定层级，比正则猜标题可靠；
- Title/Subtitle 样式视为书名，不进正文；
- 比「章」更深的标题（通常是节）作为章内小节，kind='heading'，不参与问答检索；
- 没有标题样式时退回文本模式（第X章 / Chapter N / 附录A，见 base.build_blocks）；
- **读段落内的 run**（`paragraph.runs`），把粗体/斜体/上下标/字号差异带进
  `ParsedSection.content`，阅读页因此能还原 Word 里的行内版式；
- Word 没有物理页码，段落页号为按字数的**虚拟页码**（见 base 模块说明）。

说明：**只读得到正文顶层段落**（`doc.paragraphs`）——文本框、图片里的文字
一律读不到，它们由 `skipped_notes()` 计数后写进 `ParsedBook.notes`，上传后提示用户。
表格同理，但表格在 phase 3 会被单独提取成 `kind='table'` 的段落。
"""
import io
import logging
import re
from typing import Any, Dict, List, Optional, Tuple

from docx import Document

from .base import (
    ParsedBook,
    Run,
    assemble_book,
    build_blocks,
    make_run,
    merge_runs,
    runs_payload,
    text_of_runs,
)

logger = logging.getLogger(__name__)

_HEADING_RE = re.compile(r"^Heading\s+(\d+)$", re.IGNORECASE)
#: 视为书名的 Word 样式
_TITLE_STYLES = ("title", "subtitle")
#: WordprocessingML 命名空间：直接查 XML 里的表格 / 文本框 / 图片节点
_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

#: 相对正文基准字号多大算「大一号 / 小一号」
_BIG_RATIO = 1.15
_SMALL_RATIO = 0.9


def _level_of(style_name: str) -> int:
    """Word 样式名 → 层级：Heading N = N，其余（含 Title）= 0。"""
    match = _HEADING_RE.match((style_name or "").strip())
    return min(int(match.group(1)), 9) if match else 0


def _run_style(run, body_size: Optional[float]) -> Tuple[str, ...]:
    """Word 的 run 格式 → 语义样式 token（见 base.STYLE_TOKENS）。

    上下标只认 `font.subscript` / `font.superscript`：python-docx 把它们映射到
    `<w:vertAlign>`，而 `<w:position>`（基线微调）**没有暴露成属性**，读写都得碰
    XML。教材里的上下标都用 vertAlign，所以不额外去解析 XML。
    """
    style: List[str] = []
    if run.bold:
        style.append("b")
    if run.italic:
        style.append("i")

    font = run.font
    if font.subscript:
        style.append("sub")
    elif font.superscript:
        style.append("sup")

    if "sub" not in style and "sup" not in style:
        size = font.size.pt if font.size is not None else None
        if size and body_size:
            if size >= body_size * _BIG_RATIO:
                style.append("lg")
            elif size <= body_size * _SMALL_RATIO:
                style.append("sm")
    return tuple(style)


def _size_histogram(paragraphs) -> Dict[float, int]:
    """显式设了字号的 run 的尺寸频次，用于求正文基准字号。"""
    counts: Dict[float, int] = {}
    for paragraph in paragraphs:
        for run in paragraph.runs:
            size = run.font.size.pt if run.font.size is not None else None
            if size:
                counts[size] = counts.get(size, 0) + 1
    return counts


def _body_size(paragraphs) -> Optional[float]:
    """正文基准字号 = **出现最多**的字号（并列时取较小者）。

    不能用中位数：Word 里字号只在显式设置过的地方才有值，样本很稀（可能只有
    标题的大字号和正文各一次），中位数会被大字号带偏，导致正文自己被判成 lg。
    取众数才对应「正文用什么字号排的」这个语义。
    """
    counts = _size_histogram(paragraphs)
    if not counts:
        return None
    top = max(counts.values())
    return min(size for size, count in counts.items() if count == top)


def styled_runs_of(paragraph, body_size: Optional[float]) -> Tuple[List[Run], str]:
    """读一个段落的行内片段，返回 (片段列表, 纯文本)。

    纯文本取自 `paragraph.text`（Word 自己拼好的、含超链接文字的完整文本），
    只在两者**不一致**时退回「整段无样式」：`paragraph.runs` 拿不到 `<w:hyperlink>`
    或域代码里的 run，直接用片段拼文本会**静默丢掉**那部分文字。宁可少一点样式，
    也不能少字。
    """
    runs: List[Run] = []
    for run in paragraph.runs:
        piece = make_run(run.text or "", *_run_style(run, body_size))
        if piece is not None:
            runs.append(piece)

    full = paragraph.text or ""
    # 全段都没有行内格式时不产出 content：渲染层走「整段纯文本」的回退路径，
    # 与纯文本/Markdown 教材以及升级前入库的老数据完全一致，不多一层包裹。
    if all(not run.style for run in runs):
        return [], full
    joined = text_of_runs(runs_payload(runs)["runs"]) if runs else ""
    if joined != full:
        if joined:
            logger.debug("段落含未展开的行内节点，退回纯文本：%r", full[:40])
        return [], full
    return merge_runs(runs), full


def document_rows(document) -> List[Tuple[str, str, int, Optional[Dict[str, Any]]]]:
    """读出 [(纯文本, 样式名, 层级, 富文本内容)]，跳过空段。"""
    paragraphs = document.paragraphs
    body_size = _body_size(paragraphs)
    rows: List[Tuple[str, str, int, Optional[Dict[str, Any]]]] = []
    for paragraph in paragraphs:
        text = (paragraph.text or "").strip()
        if not text:
            continue
        style = paragraph.style.name if paragraph.style is not None else "Normal"
        runs, plain = styled_runs_of(paragraph, body_size)
        rows.append((plain, style, _level_of(style), runs_payload(runs)))
    return rows


def read_rows(data: bytes) -> List[Tuple[str, int, str]]:
    """读出 [(段落文本, 层级, 样式名)]，跳过空段（沿用旧的公开形状）。"""
    return [(text, level, style) for text, style, level, _ in document_rows(Document(io.BytesIO(data)))]


def rows_of(document) -> List[Tuple[str, int, str]]:
    """从已打开的文档读出段落行（parse 与「跳过了什么」共用同一个 Document）。"""
    return [(text, level, style) for text, style, level, _ in document_rows(document)]


def skipped_notes(document) -> List[str]:
    """报告解析器**读不到**的内容，供上传后提示用。

    `doc.paragraphs` 只覆盖正文顶层段落，下面这三类里都有文字但读不到，
    也正是「上传后整本书只剩几百字」最常见的原因。

    表格这一项在 phase 3 会被单独提取成 `kind='table'` 的段落，届时这里
    要改成只报真正还读不到的（文本框、图片里的文字）；现在表格确实一个字
    都没进正文，如实报告。
    """
    def count(tag: str) -> int:
        return len(document.element.body.findall(f".//{{{_W_NS}}}{tag}"))

    found: List[str] = []
    for tag, label in (("tbl", "个表格"), ("txbxContent", "个文本框"), ("drawing", "张图片")):
        total = count(tag)
        if total:
            found.append(f"{total} {label}")
    return found


def parse_docx_bytes(
    data: bytes, default_title: str = "未命名教材", assets_dir=None
) -> ParsedBook:
    """解析 .docx 字节流。

    `assets_dir` 给出时，文档内嵌图片会落盘到该目录（阶段 3 启用）；为 None
    则纯文本解析、不产生任何文件。参数先占位，让 ingest 的调用契约现在就稳定。
    """
    document = Document(io.BytesIO(data))
    rows = document_rows(document)

    title = ""
    body: List[Tuple[str, int, Optional[Dict[str, Any]]]] = []
    for text, style, level, content in rows:
        if style.lower() in _TITLE_STYLES:
            title = title or text  # 书名只取第一次出现的 Title/Subtitle
            continue
        body.append((text, level, content))

    return assemble_book(title, build_blocks(body), default_title, notes=skipped_notes(document))
