"""生成一份“原创样例教材” .docx（版权安全，供流水线与测试使用）。

内容为自行撰写的极简中文讲义示例，不含任何受版权保护的教材内容。
用法：
    .venv\\Scripts\\python.exe scripts\\make_sample_book.py
默认输出到 data/source/sample_book.docx（data/ 已被 .gitignore 忽略）。
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document

BOOK_TITLE = "微积分简明讲义（样例）"
AUTHOR = "kb-agent 示例作者"

CHAPTERS = [
    {
        "title": "第一章 函数与极限",
        "paragraphs": [
            "函数是描述两个变量之间依赖关系的数学对象：对定义域内的每一个自变量 x，都有唯一确定的因变量 y 与之对应。",
            "极限刻画的是变量在变化过程中无限趋近于某个确定值的趋势，是微积分的基石概念。",
            "例如，当自变量 x 越来越接近 1 时，函数 f(x) = x 的取值也会越来越接近 1，我们就说 f(x) 在 x 趋向 1 时的极限为 1。",
            "理解极限需要注意：极限描述的是“趋近的过程”，而不是函数在该点是否真的有定义或取值。",
        ],
    },
    {
        "title": "第二章 导数与微分",
        "paragraphs": [
            "导数描述函数在某一点处的瞬时变化率，也就是函数曲线在该点切线的斜率。",
            "设函数 y = f(x) 在点 x0 的某个邻域内有定义，当自变量在点 x0 处取得增量 Δx 时，函数相应地取得增量 Δy = f(x0 + Δx) - f(x0)。",
            "如果当 Δx 趋近于 0 时，比值 Δy / Δx 的极限存在，那么称这个极限为函数 y = f(x) 在点 x0 处的导数。",
            "导数公式示例：f′(x0) = lim(Δx→0) [f(x0 + Δx) - f(x0)] / Δx。",
            "导数的几何意义是曲线在对应点处切线的斜率；物理意义则是该时刻的瞬时变化率。",
        ],
    },
]


def build_sample_book(path: Path) -> Path:
    """按 CHAPTERS 生成样例教材 docx，返回输出路径。"""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    doc.add_heading(BOOK_TITLE, level=0)
    doc.add_paragraph(f"作者：{AUTHOR}（本文件为流水线自测用的原创示例，可自由使用）")

    for chapter in CHAPTERS:
        doc.add_heading(chapter["title"], level=1)
        for para in chapter["paragraphs"]:
            doc.add_paragraph(para)

    doc.save(path)
    return path


def main() -> None:
    out = Path(__file__).resolve().parent.parent / "data" / "source" / "sample_book.docx"
    saved = build_sample_book(out)
    print(f"已生成样例教材：{saved}")


if __name__ == "__main__":
    sys.exit(main())
